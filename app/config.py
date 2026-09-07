from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any
import json
from openpyxl import load_workbook

from app.paths import CONFIG_DIR


@dataclass(frozen=True)
class Project:
    project_id: str
    project_name: str
    postal_code: str


@dataclass(frozen=True)
class Competitor:
    project_id: str
    asin: str
    internal_name: str = ""
    brand: str = ""
    remark: str = ""


@dataclass(frozen=True)
class Keyword:
    project_id: str
    keyword: str
    pages: int


@dataclass(frozen=True)
class Category:
    project_id: str
    category_name: str
    category_url: str
    max_rank: int


def load_settings(path: Path | None = None) -> dict[str, Any]:
    settings_path = path or CONFIG_DIR / "settings.json"
    with settings_path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def _read_rows(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"配置文件不存在: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration:
        raise ValueError(f"配置文件为空: {path}") from None
    result: list[dict[str, Any]] = []
    for row_number, values in enumerate(rows, start=2):
        row = dict(zip(headers, values))
        row["_row"] = row_number
        result.append(row)
    return result


def _enabled(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "y", "是"}


def load_projects() -> list[Project]:
    items: list[Project] = []
    for row in _read_rows(CONFIG_DIR / "projects.xlsx"):
        if not _enabled(row.get("enabled")):
            continue
        project_id = str(row.get("project_id") or "").strip()
        project_name = str(row.get("project_name") or "").strip()
        if not project_id or not project_name:
            raise ValueError(f"projects.xlsx 第 {row['_row']} 行缺少 project_id 或 project_name")
        items.append(Project(project_id, project_name, str(row.get("postal_code") or "90001").strip()))
    if len({item.project_id for item in items}) != len(items):
        raise ValueError("projects.xlsx 中存在重复的 project_id")
    return items


def load_competitors(project_id: str | None = None, include_disabled: bool = False) -> list[Competitor]:
    items: list[Competitor] = []
    for row in _read_rows(CONFIG_DIR / "competitors.xlsx"):
        if not include_disabled and not _enabled(row.get("enabled")):
            continue
        row_project = str(row.get("project_id") or "").strip()
        if not row_project:
            raise ValueError(f"competitors.xlsx 第 {row['_row']} 行缺少 project_id")
        if project_id and row_project != project_id:
            continue
        asin = str(row.get("asin") or "").strip().upper()
        if not asin:
            raise ValueError(f"competitors.xlsx 第 {row['_row']} 行缺少 asin")
        items.append(Competitor(row_project, asin, str(row.get("internal_name") or ""), str(row.get("brand") or ""), str(row.get("remark") or "")))
    return items


def load_keywords(default_pages: int = 3, project_id: str | None = None) -> list[Keyword]:
    items: list[Keyword] = []
    for row in _read_rows(CONFIG_DIR / "keywords.xlsx"):
        if not _enabled(row.get("enabled")):
            continue
        row_project = str(row.get("project_id") or "").strip()
        if not row_project:
            raise ValueError(f"keywords.xlsx 第 {row['_row']} 行缺少 project_id")
        if project_id and row_project != project_id:
            continue
        keyword = str(row.get("keyword") or "").strip()
        if not keyword:
            raise ValueError(f"keywords.xlsx 第 {row['_row']} 行缺少 keyword")
        pages = int(row.get("pages") or default_pages)
        items.append(Keyword(row_project, keyword, pages))
    return items


def load_categories(project_id: str | None = None) -> list[Category]:
    items: list[Category] = []
    for row in _read_rows(CONFIG_DIR / "categories.xlsx"):
        if not _enabled(row.get("enabled")):
            continue
        row_project = str(row.get("project_id") or "").strip()
        if not row_project:
            raise ValueError(f"categories.xlsx 第 {row['_row']} 行缺少 project_id")
        if project_id and row_project != project_id:
            continue
        name = str(row.get("category_name") or "").strip()
        url = str(row.get("category_url") or "").strip()
        if not name or not url:
            raise ValueError(f"categories.xlsx 第 {row['_row']} 行缺少 category_name 或 category_url")
        items.append(Category(row_project, name, url, int(row.get("max_rank") or 100)))
    return items
