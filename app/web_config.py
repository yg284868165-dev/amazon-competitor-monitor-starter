from __future__ import annotations

import json
import os
import re
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from app.candidates import load_rules, save_rules
from app.paths import CONFIG_DIR
from app.storage.database import Database

CONFIG_SCHEMAS = {
    "projects": ["enabled", "project_id", "project_name", "postal_code"],
    "competitors": ["enabled", "project_id", "asin", "internal_name", "brand", "remark"],
    "keywords": ["enabled", "project_id", "keyword", "pages"],
    "categories": ["enabled", "project_id", "category_name", "category_url", "max_rank"],
}


def read_table(kind: str) -> list[dict[str, Any]]:
    headers = CONFIG_SCHEMAS[kind]
    workbook = load_workbook(CONFIG_DIR / f"{kind}.xlsx", read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    actual = [str(value or "") for value in rows[0]]
    indexes = {name: actual.index(name) for name in headers if name in actual}
    return [
        {name: values[indexes[name]] if name in indexes and indexes[name] < len(values) else "" for name in headers}
        for values in rows[1:]
        if any(value not in (None, "") for value in values)
    ]


def _validate(kind: str, rows: list[dict[str, Any]], project_ids_override: set[str] | None = None) -> list[dict[str, Any]]:
    if not isinstance(rows, list):
        raise ValueError("配置数据必须是行列表")
    cleaned = []
    project_ids = (
        project_ids_override
        if project_ids_override is not None
        else {str(row.get("project_id") or "").strip() for row in read_table("projects")}
    ) if kind != "projects" else set()
    seen = set()
    for index, row in enumerate(rows, start=2):
        normalized = {name: row.get(name, "") for name in CONFIG_SCHEMAS[kind]}
        normalized["enabled"] = 1 if str(normalized["enabled"]).lower() in {"1", "true", "yes", "是"} or normalized["enabled"] is True else 0
        project_id = str(normalized.get("project_id") or "").strip()
        if kind != "projects" and not project_id:
            raise ValueError(f"第 {index} 行请选择产品项目")
        if not re.fullmatch(r"[a-z0-9_-]+", project_id):
            raise ValueError(f"第 {index} 行 project_id 只能使用小写字母、数字、下划线和短横线")
        if kind != "projects" and project_id not in project_ids:
            raise ValueError(f"第 {index} 行 project_id={project_id} 不存在于产品项目中")
        if kind == "projects":
            if not str(normalized.get("project_name") or "").strip():
                raise ValueError(f"第 {index} 行缺少项目名称")
            postal = str(normalized.get("postal_code") or "").strip()
            if not re.fullmatch(r"\d{5}", postal):
                raise ValueError(f"第 {index} 行美国邮编必须是 5 位数字")
            unique = project_id
        elif kind == "competitors":
            asin = str(normalized.get("asin") or "").strip().upper()
            if not re.fullmatch(r"[A-Z0-9]{10}", asin):
                raise ValueError(f"第 {index} 行 ASIN 必须是 10 位字母或数字")
            normalized["asin"] = asin
            unique = (project_id, asin)
        elif kind == "keywords":
            if not str(normalized.get("keyword") or "").strip():
                raise ValueError(f"第 {index} 行关键词不能为空")
            pages = int(normalized.get("pages") or 3)
            if not 1 <= pages <= 10:
                raise ValueError(f"第 {index} 行 pages 必须在 1–10 之间")
            normalized["pages"] = pages
            unique = (project_id, str(normalized["keyword"]).strip().lower())
        else:
            if not str(normalized.get("category_name") or "").strip():
                raise ValueError(f"第 {index} 行类目名称不能为空")
            url = str(normalized.get("category_url") or "").strip()
            if not url.startswith("https://www.amazon.com/"):
                raise ValueError(f"第 {index} 行必须填写 amazon.com 的 Best Sellers 链接")
            normalized["max_rank"] = int(normalized.get("max_rank") or 100)
            unique = (project_id, url)
        if unique in seen:
            raise ValueError(f"第 {index} 行与前面的配置重复")
        seen.add(unique)
        cleaned.append(normalized)
    return cleaned


def _write_clean_table(kind: str, cleaned: list[dict[str, Any]]) -> None:
    headers = CONFIG_SCHEMAS[kind]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "配置"
    sheet.append(headers)
    for row in cleaned:
        sheet.append([row.get(name, "") for name in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 80)
    target = CONFIG_DIR / f"{kind}.xlsx"
    descriptor, temp_name = tempfile.mkstemp(prefix=f".{kind}_", suffix=".xlsx", dir=CONFIG_DIR)
    os.close(descriptor)
    try:
        workbook.save(temp_name)
        os.replace(temp_name, target)
    finally:
        Path(temp_name).unlink(missing_ok=True)


def write_table(kind: str, rows: list[dict[str, Any]]) -> None:
    if kind != "projects":
        _write_clean_table(kind, _validate(kind, rows))
        return

    current_projects = read_table("projects")
    old_ids = {str(row.get("project_id") or "").strip() for row in current_projects}
    cleaned_projects = _validate("projects", rows)
    new_ids = {str(row["project_id"]) for row in cleaned_projects}
    renames: dict[str, str] = {}
    submitted_originals: set[str] = set()
    for raw, cleaned in zip(rows, cleaned_projects):
        new_id = str(cleaned["project_id"])
        original = str(raw.get("_original_project_id") or "").strip()
        if not original and new_id in old_ids:
            original = new_id
        if original in old_ids:
            submitted_originals.add(original)
            if original != new_id:
                renames[original] = new_id

    removed = old_ids - submitted_originals
    raw_child_tables: dict[str, list[dict[str, Any]]] = {}
    referenced_removed: set[str] = set()
    for child_kind in ("competitors", "keywords", "categories"):
        child_rows = read_table(child_kind)
        for row in child_rows:
            old_project_id = str(row.get("project_id") or "")
            if old_project_id in removed:
                referenced_removed.add(old_project_id)
            if old_project_id in renames:
                row["project_id"] = renames[old_project_id]
        raw_child_tables[child_kind] = child_rows
    if referenced_removed:
        raise ValueError(f"以下项目仍被其他配置引用，不能删除: {', '.join(sorted(referenced_removed))}")

    child_tables = {
        child_kind: _validate(child_kind, child_rows, new_ids)
        for child_kind, child_rows in raw_child_tables.items()
    }
    for child_kind, child_rows in child_tables.items():
        _write_clean_table(child_kind, child_rows)
    _write_clean_table("projects", cleaned_projects)
    database = Database()
    for old_project_id, new_project_id in renames.items():
        database.rename_project(old_project_id, new_project_id)
    rules = load_rules()
    rules_changed = False
    for old_project_id, new_project_id in renames.items():
        if old_project_id in rules:
            rules[new_project_id] = rules.pop(old_project_id)
            rules_changed = True
    if rules_changed:
        save_rules(rules)
    schedule = read_schedule()
    scheduled_project = str(schedule.get("project") or "all")
    if scheduled_project in renames:
        schedule["project"] = renames[scheduled_project]
        write_schedule(schedule)


def parse_asin_batch(value: str) -> tuple[list[str], list[str]]:
    tokens = [token.strip().upper() for token in re.split(r"[\s,;，；]+", value or "") if token.strip()]
    valid: list[str] = []
    invalid: list[str] = []
    for token in tokens:
        target = valid if re.fullmatch(r"[A-Z0-9]{10}", token) else invalid
        if token not in target:
            target.append(token)
    return valid, invalid


def bulk_add_competitors(project_id: str, value: str, remark: str = "") -> dict[str, Any]:
    projects = {str(row["project_id"]) for row in read_table("projects")}
    if project_id not in projects:
        raise ValueError("请选择有效的产品项目")
    valid, invalid = parse_asin_batch(value)
    if not valid and not invalid:
        raise ValueError("请至少粘贴一个 ASIN")
    rows = read_table("competitors")
    existing = {str(row.get("asin") or "").upper() for row in rows if str(row.get("project_id")) == project_id}
    duplicates = [asin for asin in valid if asin in existing]
    added = [asin for asin in valid if asin not in existing]
    for asin in added:
        rows.append({
            "enabled": 1, "project_id": project_id, "asin": asin,
            "internal_name": "", "brand": "", "remark": remark.strip(),
        })
    if added:
        write_table("competitors", rows)
    return {"added": added, "duplicates": duplicates, "invalid": invalid}


def parse_keyword_batch(value: str) -> list[str]:
    tokens = [token.strip() for token in re.split(r"[\r\n,;，；]+", value or "") if token.strip()]
    unique: list[str] = []
    seen: set[str] = set()
    for token in tokens:
        normalized = token.casefold()
        if normalized not in seen:
            seen.add(normalized)
            unique.append(token)
    return unique


def bulk_add_keywords(project_id: str, value: str, pages: int = 3) -> dict[str, Any]:
    projects = {str(row["project_id"]) for row in read_table("projects")}
    if project_id not in projects:
        raise ValueError("请选择有效的产品项目")
    try:
        pages = int(pages)
    except (TypeError, ValueError) as exc:
        raise ValueError("搜索页数必须是 1–10 的整数") from exc
    if not 1 <= pages <= 10:
        raise ValueError("搜索页数必须在 1–10 之间")
    keywords = parse_keyword_batch(value)
    if not keywords:
        raise ValueError("请至少粘贴一个关键词")
    rows = read_table("keywords")
    existing = {
        str(row.get("keyword") or "").strip().casefold()
        for row in rows if str(row.get("project_id")) == project_id
    }
    duplicates = [keyword for keyword in keywords if keyword.casefold() in existing]
    added = [keyword for keyword in keywords if keyword.casefold() not in existing]
    for keyword in added:
        rows.append({"enabled": 1, "project_id": project_id, "keyword": keyword, "pages": pages})
    if added:
        write_table("keywords", rows)
    return {"added": added, "duplicates": duplicates}


def read_schedule() -> dict:
    with (CONFIG_DIR / "schedule.json").open("r", encoding="utf-8") as handle:
        return json.load(handle)


def write_schedule(value: dict[str, Any]) -> None:
    enabled = bool(value.get("enabled"))
    times = list(dict.fromkeys(str(item).strip() for item in value.get("times", [])))
    for item in times:
        if not re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", item):
            raise ValueError(f"无效时间: {item}")
    if not times:
        raise ValueError("至少保留一个执行时间")
    task = str(value.get("task") or "all")
    project = str(value.get("project") or "all")
    if task not in {"product", "search", "bestseller", "all"}:
        raise ValueError("无效任务类型")
    valid_projects = {row["project_id"] for row in read_table("projects")}
    if project != "all" and project not in valid_projects:
        raise ValueError("定时任务指定的产品项目不存在")
    target = CONFIG_DIR / "schedule.json"
    descriptor, temp_name = tempfile.mkstemp(prefix=".schedule_", suffix=".json", dir=CONFIG_DIR)
    os.close(descriptor)
    try:
        wake_enabled = bool(value.get("wake_enabled"))
        wake_lead_minutes = int(value.get("wake_lead_minutes") or 5)
        if not 1 <= wake_lead_minutes <= 30:
            raise ValueError("自动唤醒提前量必须在 1–30 分钟之间")
        Path(temp_name).write_text(json.dumps({"enabled": enabled, "times": times, "task": task, "project": project, "wake_enabled": wake_enabled, "wake_lead_minutes": wake_lead_minutes}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        os.replace(temp_name, target)
    finally:
        Path(temp_name).unlink(missing_ok=True)
