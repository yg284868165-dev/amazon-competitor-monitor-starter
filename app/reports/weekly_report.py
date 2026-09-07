from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook

from app.ad_patterns import (
    build_daily_ad_paths, collect_ad_observations, summarize_ad_insights,
)
from app.config import load_projects
from app.notifications.serverchan import _weekly_rating_changes, _weekly_recap_items
from app.presentation import (
    FIELD_LABELS, SOURCE_LABELS, build_product_identities, format_change_summary,
)
from app.reports.daily_report import _add_link, _style
from app.storage.database import Database
from app.trends import METRIC_LABELS, analyze_all_weekly_trends


RECAP_HEADERS = [
    "周报周期", "项目", "商品", "采集日期", "截至周末状态", "变化项目", "变化摘要",
    "原值", "新值", "采集时间", "ASIN", "商品链接", "来源", "关键词/类目",
]

TREND_HEADERS = [
    "周报周期", "项目", "商品", "周报项目", "类目/关键词",
    "近1周", "近2周", "近4周", "备注", "ASIN", "商品链接",
]

AD_INSIGHT_HEADERS = [
    "周报周期", "项目", "商品", "ASIN", "关键词", "观察事实", "可能解释",
    "可信度", "各时段统计", "商品链接",
]

AD_PATH_HEADERS = [
    "近28天周期", "美国西部日期", "项目", "商品", "ASIN", "关键词",
    "分时轨迹（美国西部时间）", "当日最好广告位", "当日最差广告位",
    "观察到广告次数", "成功采集次数", "全部采集次数", "商品链接",
]

AD_DETAIL_HEADERS = [
    "项目", "商品", "ASIN", "关键词", "美国西部时间", "北京时间", "采集结果",
    "广告位", "对应时段", "批次类型", "批次ID", "商品链接",
]


def build_weekly_excel_report(week_end: date, db: Database | None = None) -> BytesIO:
    if week_end.weekday() != 6:
        raise ValueError("周报统计截止日期必须是周日")
    database = db or Database()
    week_start = week_end - timedelta(days=6)
    period = f"{week_start.isoformat()} 至 {week_end.isoformat()}"
    recap_items = _weekly_recap_items(database, week_end)
    ratings = _weekly_rating_changes(database, week_end)
    trends = analyze_all_weekly_trends(database, week_end)
    ad_start = week_end - timedelta(days=27)
    ad_observations = collect_ad_observations(database, ad_start, week_end)
    ad_insights = summarize_ad_insights(ad_observations, week_start, week_end)
    ad_paths = build_daily_ad_paths(ad_observations)
    projects = {item.project_id: item.project_name for item in load_projects()}
    project_ids = list(dict.fromkeys([
        *projects,
        *(row["project_id"] for row in recap_items),
        *(row["project_id"] for row in ratings),
        *(row["project_id"] for row in trends),
        *(row["project_id"] for row in ad_insights),
        *(row["project_id"] for row in ad_observations),
    ]))
    identities = {
        project_id: build_product_identities(database, project_id)
        for project_id in project_ids
    }

    workbook = Workbook()
    recap_sheet = workbook.active
    recap_sheet.title = "本周重点复盘"
    recap_sheet.append(RECAP_HEADERS)
    trend_sheet = workbook.create_sheet("评价与排名趋势")
    trend_sheet.append(TREND_HEADERS)
    ad_insight_sheet = workbook.create_sheet("广告投放观察")
    ad_insight_sheet.append(AD_INSIGHT_HEADERS)
    ad_path_sheet = workbook.create_sheet("广告分时轨迹")
    ad_path_sheet.append(AD_PATH_HEADERS)
    ad_detail_sheet = workbook.create_sheet("广告采集明细")
    ad_detail_sheet.append(AD_DETAIL_HEADERS)

    for project_id in project_ids:
        project_name = projects.get(project_id, project_id)
        product_names = identities.get(project_id, {})
        project_recap = [row for row in recap_items if row["project_id"] == project_id]
        by_product: dict[str, list[dict]] = {}
        for row in project_recap:
            key = str(row.get("asin") or row.get("keyword") or row.get("category_name") or "其他变化")
            by_product.setdefault(key, []).append(row)
        for product_key, product_items in by_product.items():
            for row in product_items:
                asin = str(row.get("asin") or "")
                product_url = f"https://www.amazon.com/dp/{asin}" if asin else ""
                context = row.get("keyword") or row.get("category_name") or ""
                recap_sheet.append([
                    period, project_name, product_names.get(asin, asin or product_key),
                    row.get("report_date"), row.get("weekly_state"),
                    FIELD_LABELS.get(row.get("field_name"), row.get("field_name") or "变化"),
                    format_change_summary(row), row.get("old_value"), row.get("new_value"),
                    row.get("event_time"), asin, product_url,
                    SOURCE_LABELS.get(row.get("source_type"), row.get("source_type")), context,
                ])
                _add_link(recap_sheet.cell(recap_sheet.max_row, 8), row.get("old_value"))
                _add_link(recap_sheet.cell(recap_sheet.max_row, 9), row.get("new_value"))
                _add_link(recap_sheet.cell(recap_sheet.max_row, 12), product_url)
        for row in (item for item in ratings if item["project_id"] == project_id):
            asin = str(row.get("asin") or "")
            direction = "增加" if row["change_count"] > 0 else "减少"
            summary = (
                f"{row['old_count']} → {row['new_count']}"
                f"（{direction}{abs(row['change_count'])}条）"
            )
            _append_row(
                trend_sheet, period, project_name, product_names.get(asin, asin or "—"),
                "评价数量变化", "", summary, "", "", "", asin,
            )
        for row in (item for item in trends if item["project_id"] == project_id):
            asin = str(row.get("asin") or "")
            summaries = list(row.get("summaries") or [])
            summaries.extend([""] * (3 - len(summaries)))
            context = row.get("category_name") or row.get("keyword") or ""
            _append_row(
                trend_sheet, period, project_name, row.get("product") or product_names.get(asin, asin or "—"),
                METRIC_LABELS.get(row.get("metric"), row.get("metric") or "排名趋势"),
                context, summaries[0], summaries[1], summaries[2], "", asin,
            )

    for row in ad_insights:
        asin = str(row.get("asin") or "")
        product_url = f"https://www.amazon.com/dp/{asin}" if asin else ""
        stats = []
        for slot, value in row.get("slot_statistics", {}).items():
            rank = f"，中位第{value['median_rank']}位" if value.get("median_rank") is not None else ""
            stats.append(f"美西{slot[:2]}时 {value['observed']}/{value['successful']}次{rank}")
        ad_insight_sheet.append([
            period, row.get("project_name") or projects.get(row["project_id"], row["project_id"]),
            row.get("product"), asin, row.get("keyword"), row.get("observation"),
            row.get("inference"), row.get("confidence"), "；".join(stats), product_url,
        ])
        _add_link(ad_insight_sheet.cell(ad_insight_sheet.max_row, 10), product_url)

    ad_period = f"{ad_start.isoformat()} 至 {week_end.isoformat()}（美国西部日期）"
    for row in ad_paths:
        asin = str(row.get("asin") or "")
        product_url = f"https://www.amazon.com/dp/{asin}" if asin else ""
        ad_path_sheet.append([
            ad_period, row.get("pacific_date"), row.get("project_name"), row.get("product"),
            asin, row.get("keyword"), row.get("path"), row.get("best_rank"),
            row.get("worst_rank"), row.get("observed_count"), row.get("successful_count"),
            row.get("sample_count"), product_url,
        ])
        _add_link(ad_path_sheet.cell(ad_path_sheet.max_row, 13), product_url)

    result_labels = {"observed": "观察到广告", "not_observed": "未观察到广告", "failed": "采集失败"}
    for row in ad_observations:
        asin = str(row.get("asin") or "")
        product_url = f"https://www.amazon.com/dp/{asin}" if asin else ""
        ad_detail_sheet.append([
            row.get("project_name"), row.get("product"), asin, row.get("keyword"),
            row.get("pacific_datetime"), row.get("china_datetime"), result_labels.get(row.get("status"), row.get("status")),
            row.get("ad_rank"), row.get("slot"), row.get("run_type"), row.get("run_id"), product_url,
        ])
        _add_link(ad_detail_sheet.cell(ad_detail_sheet.max_row, 12), product_url)

    if recap_sheet.max_row == 1:
        recap_sheet.append([period, "全部项目", "上周未发现日报重点变化"])
    if trend_sheet.max_row == 1:
        trend_sheet.append([period, "全部项目", "上周未发现评价数量净变化或明显排名趋势"])
    if ad_insight_sheet.max_row == 1:
        ad_insight_sheet.append([period, "全部项目", "上周未发现可信度达到中或高的广告投放规律"])
    if ad_path_sheet.max_row == 1:
        ad_path_sheet.append([ad_period, "", "全部项目", "近28天没有可用的广告采集记录"])
    if ad_detail_sheet.max_row == 1:
        ad_detail_sheet.append(["全部项目", "近28天没有可用的广告采集记录"])
    _style(recap_sheet)
    _style(trend_sheet)
    _style(ad_insight_sheet)
    _style(ad_path_sheet)
    _style(ad_detail_sheet)
    payload = BytesIO()
    workbook.save(payload)
    payload.seek(0)
    return payload


def _append_row(
    sheet, period: str, project_name: str, product: str, item: str, context: str,
    week_1: str, week_2: str, week_4: str, note: str, asin: str,
) -> None:
    product_url = f"https://www.amazon.com/dp/{asin}" if asin else ""
    sheet.append([
        period, project_name, product, item, context,
        week_1, week_2, week_4, note, asin, product_url,
    ])
    _add_link(sheet.cell(sheet.max_row, 11), product_url)
