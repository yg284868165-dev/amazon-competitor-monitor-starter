from __future__ import annotations

import json
from io import BytesIO
from typing import Any

from openpyxl import Workbook

from app.asin_history import build_asin_history
from app.config import load_projects
from app.reports.daily_report import _add_link, _style
from app.storage.database import Database


POINT_LABELS = {
    "ranked": lambda point: f"第{point['value']}名",
    "unranked": lambda point: "未进入监控页数",
    "not_observed": lambda point: "该次未显示类目排名",
    "pending": lambda point: "待积累",
    "no_data": lambda point: "无有效采样",
}
AD_STATUS_LABELS = {
    "observed": "观察到广告",
    "not_observed": "未观察到广告",
    "collection_failed": "采集失败",
}
LISTING_STATUS_LABELS = {
    "active": "正常",
    "dog": "页面变狗",
    "removed": "商品下架",
}


def _text(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, list):
        return "\n".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _point_text(point: dict[str, Any]) -> str:
    formatter = POINT_LABELS.get(str(point.get("status")))
    result = formatter(point) if formatter else str(point.get("status") or "")
    observed_at = point.get("observed_at")
    return f"{result}\n采集于 {observed_at}" if observed_at else result


def build_asin_history_excel(
    project_id: str, asin: str, days: int = 30, db: Database | None = None,
) -> BytesIO:
    database = db or Database()
    history = build_asin_history(database, project_id, asin, days)
    project_names = {item.project_id: item.project_name for item in load_projects()}
    project_name = project_names.get(project_id, project_id)
    product_url = history["amazon_url"]

    workbook = Workbook()
    overview = workbook.active
    overview.title = "动作摘要"
    overview.append(["项目", "内容", "说明"])
    overview_rows = [
        ("产品项目", project_name, project_id),
        ("商品", history["identity"], history["asin"]),
        ("查看范围", f"近{days}天", f"{history['range']['from']} 至 {history['range']['to']}"),
        ("有效商品快照", history["range"]["sample_count"], "实际有数据的采集次数"),
        ("重要动作节点", history["summary"]["important_nodes"], "价格促销、Listing内容或销售状态"),
        ("价格与促销", history["summary"]["price_promo"], "按采集批次合并"),
        ("Listing内容", history["summary"]["listing"], "按采集批次合并"),
        ("销售状态", history["summary"]["operations"], "按采集批次合并"),
        ("Amazon标识", history["summary"]["platform"], "按采集批次合并"),
        ("商品链接", product_url, "可点击跳转"),
        ("时间含义", "采集时间", "表示程序检测到变化的时间，不代表竞对实际修改时间"),
    ]
    for row in overview_rows:
        overview.append(row)
    _add_link(overview.cell(overview.max_row - 1, 2), product_url)

    timeline = workbook.create_sheet("动作时间线")
    timeline.append(["采集时间", "动作节点", "分类", "记录类型", "变化项目/观察指标", "变化摘要/类目关键词", "原值/动作前", "新值/1天后", "3天后", "7天后", "批次ID", "ASIN", "商品链接"])
    for action in history["actions"]:
        for item in action["items"]:
            timeline.append([
                action["observed_at"], action["title"], "、".join(action["category_labels"]),
                "竞对动作", item["label"], item["summary"], _text(item["old"]), _text(item["new"]),
                "", "", action["run_id"], history["asin"], product_url,
            ])
            _add_link(timeline.cell(timeline.max_row, 13), product_url)
        for impact in action.get("impacts", []):
            timeline.append([
                action["observed_at"], action["title"], "、".join(action["category_labels"]),
                "后续排名对照", impact["metric_label"], impact["context"],
                _point_text(impact["points"]["before"]), _point_text(impact["points"]["day_1"]),
                _point_text(impact["points"]["day_3"]), _point_text(impact["points"]["day_7"]),
                action["run_id"], history["asin"], product_url,
            ])
            _add_link(timeline.cell(timeline.max_row, 13), product_url)
    if timeline.max_row == 1:
        timeline.append(["", "所选时间内没有检测到动作"])

    listing = workbook.create_sheet("Listing版本对比")
    listing.append(["采集时间", "变化项目", "变化前", "变化后", "变化摘要", "批次ID", "商品链接"])
    for change in history["listing_changes"]:
        item = change["item"]
        listing.append([
            change["observed_at"], item["label"], _text(item["old"]), _text(item["new"]),
            item["summary"], change["run_id"], product_url,
        ])
        _add_link(listing.cell(listing.max_row, 7), product_url)
        if item["display_type"] == "images":
            old_images, new_images = item["old"], item["new"]
            if old_images:
                _add_link(listing.cell(listing.max_row, 3), old_images[0])
            if new_images:
                _add_link(listing.cell(listing.max_row, 4), new_images[0])
    if listing.max_row == 1:
        listing.append(["", "所选时间内没有Listing内容变化"])

    product = workbook.create_sheet("日内商品状态")
    product.append(["采集时间", "价格", "优惠券", "促销", "企业价", "库存状态", "跟卖数量", "购物车卖家", "发货方", "页面状态", "商品链接"])
    for row in history["intraday"]["product"]:
        product.append([
            row["collected_at"], row["price"], row["coupon"], row["deal"], row["business_price"],
            row["availability"], row["offer_count"], row["featured_seller"], row["ships_from"],
            LISTING_STATUS_LABELS.get(row["listing_status"], row["listing_status"]), product_url,
        ])
        _add_link(product.cell(product.max_row, 11), product_url)
    if product.max_row == 1:
        product.append(["", "所选时间内没有有效商品快照"])

    ads = workbook.create_sheet("广告分时证据")
    ads.append(["采集时间", "关键词", "采集结果", "广告位", "商品链接"])
    for row in history["intraday"]["ads"]:
        ads.append([
            row["collected_at"], row["keyword"], AD_STATUS_LABELS.get(row["status"], row["status"]),
            f"第{row['ad_rank']}位" if row["ad_rank"] is not None else "", product_url,
        ])
        _add_link(ads.cell(ads.max_row, 5), product_url)
    if ads.max_row == 1:
        ads.append(["", "所选时间内没有关键词采集记录"])

    for sheet in workbook.worksheets:
        _style(sheet)
    payload = BytesIO()
    workbook.save(payload)
    payload.seek(0)
    return payload
