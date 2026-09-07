from __future__ import annotations

import json
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.candidates import mark_expired_candidates
from app.config import Project, load_competitors, load_keywords
from app.paths import REPORT_DIR, legacy_change_report_filename, report_filename
from app.presentation import (
    FIELD_LABELS, SOURCE_LABELS, build_product_identities,
    format_change_summary, format_rating_breakdown,
)
from app.storage.database import Database
from app.trends import analyze_project_trends

LISTING_STATUS_LABELS = {
    "active": "正常在售",
    "dog": "页面变狗",
    "removed": "商品下架",
}

def _style(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, column in enumerate(sheet.columns, 1):
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
        sheet.column_dimensions[get_column_letter(index)].width = max(width, 10)


def _json_list(value: str | None) -> list:
    if not value:
        return []
    try:
        parsed = json.loads(value)
        return parsed if isinstance(parsed, list) else []
    except (json.JSONDecodeError, TypeError):
        return []


def _format_bsr(value: str | None, main_bsr: int | None = None) -> str:
    lines = []
    for item in _json_list(value):
        category = str(item.get("category") or "").strip().strip('"')
        rank = item.get("rank")
        if category and rank is not None and rank != main_bsr:
            lines.append(f"{category}：{rank}")
    return "\n".join(dict.fromkeys(lines))


def _format_main_bsr(value: str | None, main_bsr: int | None) -> str:
    for item in _json_list(value):
        if item.get("rank") == main_bsr:
            category = str(item.get("category") or "").strip().strip('"')
            if category:
                return f"{category}：{main_bsr}"
    return str(main_bsr) if main_bsr is not None else ""


def _format_bullets(value: str | None) -> str:
    return "\n".join(f"• {text}" for text in _json_list(value))


def _format_images(value: str | None) -> str:
    return "\n".join(_json_list(value))


def _format_high_return_rate(value) -> str:
    if value is None:
        return ""
    return "是" if bool(value) else "否"


def _format_listing_status(value: str | None) -> str:
    return LISTING_STATUS_LABELS.get(value or "active", value or "正常在售")


def _latest_products(db: Database, project_id: str) -> dict[str, dict]:
    active_rows = db.fetchall("""
        SELECT p.* FROM product_snapshots p
        JOIN (
          SELECT asin, MAX(collected_at) AS latest_at
          FROM product_snapshots WHERE success=1 AND project_id=? GROUP BY asin
        ) latest ON latest.asin=p.asin AND latest.latest_at=p.collected_at
        WHERE p.success=1 AND p.project_id=?
    """, (project_id, project_id))
    result = {row["asin"]: row for row in active_rows}
    latest_rows = db.fetchall("""
        SELECT p.* FROM product_snapshots p
        JOIN (
          SELECT asin, MAX(collected_at) AS latest_at
          FROM product_snapshots WHERE project_id=?
          AND (success=1 OR listing_status IN ('dog','removed')) GROUP BY asin
        ) latest ON latest.asin=p.asin AND latest.latest_at=p.collected_at
        WHERE p.project_id=? AND (p.success=1 OR p.listing_status IN ('dog','removed'))
    """, (project_id, project_id))
    for latest in latest_rows:
        if latest.get("listing_status") in {"dog", "removed"} and latest["asin"] in result:
            merged = dict(result[latest["asin"]])
            for field in ("listing_status", "listing_status_detail", "source_url", "collected_at", "run_id"):
                merged[field] = latest.get(field)
            result[latest["asin"]] = merged
        else:
            result[latest["asin"]] = latest
    return result


def _latest_keyword_rows(db: Database, project_id: str, keyword: str) -> list[dict]:
    runs = db.fetchall("SELECT run_id FROM search_snapshots WHERE project_id=? AND keyword=? ORDER BY collected_at DESC LIMIT 1", (project_id, keyword))
    return db.fetchall("SELECT * FROM search_snapshots WHERE project_id=? AND run_id=? AND keyword=?", (project_id, runs[0]["run_id"], keyword)) if runs else []


def _change_rows(db: Database, run_id: str, project_id: str) -> list[dict]:
    today = date.today().isoformat()
    return db.fetchall("""
        SELECT * FROM change_events WHERE project_id=? AND substr(event_time,1,10)=?
        AND (source_type<>'bestseller' OR event_type IN ('entered_bestseller','new_competitor_found'))
        AND source_type<>'search' AND COALESCE(field_name,'')<>'main_bsr'
        ORDER BY CASE severity WHEN 'high' THEN 1 WHEN 'medium' THEN 2 ELSE 3 END, event_time
    """, (project_id, today))


CHANGE_HEADERS = ["变化项目", "商品", "变化摘要", "当前值", "原值", "变化量", "时间", "来源", "ASIN", "商品链接", "关键词", "类目"]


def _append_changes(sheet, changes: list[dict], identities: dict[str, str]) -> None:
    if not changes:
        sheet.append(["本次未发现变化"])
        return
    for row in changes:
        asin = str(row.get("asin") or "")
        product_url = f"https://www.amazon.com/dp/{asin}" if asin else ""
        sheet.append([
            FIELD_LABELS.get(row["field_name"], row["field_name"]), identities.get(asin, asin or "—"),
            format_change_summary(row), row["new_value"], row["old_value"], row["change_value"],
            row["event_time"], SOURCE_LABELS.get(row["source_type"], row["source_type"]),
            asin, product_url, row["keyword"], row["category_name"],
        ])
        if product_url:
            link_cell = sheet.cell(row=sheet.max_row, column=10)
            link_cell.hyperlink = product_url
            link_cell.style = "Hyperlink"


def generate_report(db: Database, run_id: str, project: Project) -> str:
    project_id = project.project_id
    mark_expired_candidates(db, project_id)
    competitors = load_competitors(project_id)
    keywords = load_keywords(project_id=project_id)
    target_asins = [item.asin for item in competitors]
    products = _latest_products(db, project_id)
    changes = _change_rows(db, run_id, project_id)
    identities = build_product_identities(db, project_id)

    workbook = Workbook()
    workbook.remove(workbook.active)
    change_sheet = workbook.create_sheet("今日重点")
    trend_sheet = workbook.create_sheet("趋势观察")
    new_asin_sheet = workbook.create_sheet("新品雷达")
    current_sheet = workbook.create_sheet("竞品当前状态")
    error_sheet = workbook.create_sheet("采集异常")
    daily_sheet = workbook.create_sheet("今日汇总")
    daily_keyword_sheet = workbook.create_sheet("今日关键词汇总")
    history_sheet = workbook.create_sheet("竞品历史")
    keyword_history_sheet = workbook.create_sheet("关键词排名历史")
    bestseller_sheet = workbook.create_sheet("BSR前100")
    run_sheet = workbook.create_sheet("运行记录")

    change_sheet.append(CHANGE_HEADERS)
    _append_changes(change_sheet, changes, identities)

    today = date.today().isoformat()
    daily_sheet.append([
        "日期", "商品", "ASIN", "采集次数", "首次采集", "末次采集", "首次价格", "最新价格",
        "最低价格", "最高价格", "当日Coupon", "当日Deal", "当日企业价", "首次大类BSR", "最新大类BSR",
        "最佳大类BSR", "最差大类BSR", "Rating首次", "Rating最新", "Rating新增",
        "最低星级", "最高星级", "不可售次数",
    ])
    daily_products = db.fetchall(
        "SELECT * FROM product_snapshots WHERE project_id=? AND success=1 AND substr(collected_at,1,10)=? ORDER BY collected_at",
        (project_id, today),
    )
    for asin in target_asins:
        rows = [row for row in daily_products if row["asin"] == asin]
        if not rows:
            continue
        prices = [row["current_price"] for row in rows if row["current_price"] is not None]
        ranks = [row["main_bsr"] for row in rows if row["main_bsr"] is not None]
        ratings = [row["rating_count"] for row in rows if row["rating_count"] is not None]
        stars = [row["rating_value"] for row in rows if row["rating_value"] is not None]
        coupons = list(dict.fromkeys(row["coupon_text"] for row in rows if row["coupon_text"]))
        deals = list(dict.fromkeys(row["deal_text"] for row in rows if row["deal_text"]))
        business_prices = list(dict.fromkeys(row.get("business_price_text") for row in rows if row.get("business_price_text")))
        unavailable = sum(1 for row in rows if row["availability"] and "in stock" not in row["availability"].lower())
        daily_sheet.append([
            today, identities.get(asin, asin), asin, len(rows), rows[0]["collected_at"], rows[-1]["collected_at"],
            prices[0] if prices else None, prices[-1] if prices else None,
            min(prices) if prices else None, max(prices) if prices else None,
            "\n".join(coupons), "\n".join(deals), "\n".join(business_prices), ranks[0] if ranks else None,
            ranks[-1] if ranks else None, min(ranks) if ranks else None, max(ranks) if ranks else None,
            ratings[0] if ratings else None, ratings[-1] if ratings else None,
            (ratings[-1] - ratings[0]) if ratings else None,
            min(stars) if stars else None, max(stars) if stars else None, unavailable,
        ])

    daily_keyword_sheet.append(["日期", "商品", "关键词", "ASIN", "采集批次", "最佳自然位", "最新自然位", "最佳广告位", "最新广告位"])
    for item in keywords:
        rows = db.fetchall(
            "SELECT * FROM search_snapshots WHERE project_id=? AND keyword=? AND substr(collected_at,1,10)=? ORDER BY collected_at",
            (project_id, item.keyword, today),
        )
        for asin in target_asins:
            matches = [row for row in rows if row["asin"] == asin]
            if not matches:
                continue
            by_run: dict[str, list[dict]] = {}
            for row in matches:
                by_run.setdefault(row["run_id"], []).append(row)
            samples = []
            for run_rows in by_run.values():
                timestamp = max(row["collected_at"] for row in run_rows)
                organic = min((row["organic_rank"] for row in run_rows if row["organic_rank"] is not None), default=None)
                ad = min((row["ad_rank"] for row in run_rows if row["ad_rank"] is not None), default=None)
                samples.append((timestamp, organic, ad))
            samples.sort(key=lambda item: item[0])
            organic_values = [sample[1] for sample in samples if sample[1] is not None]
            ad_values = [sample[2] for sample in samples if sample[2] is not None]
            daily_keyword_sheet.append([
                today, identities.get(asin, asin), item.keyword, asin, len(samples), min(organic_values) if organic_values else None,
                samples[-1][1], min(ad_values) if ad_values else None, samples[-1][2],
            ])

    trend_sheet.append(["周期", "商品", "趋势项目", "类目/关键词", "起始代表排名", "当前代表排名", "方向", "变化名次", "有效天数", "说明"])
    for trend in analyze_project_trends(db, project_id, date.today()):
        trend_sheet.append([
            f"{trend['period']}天", trend["product"], trend["metric_label"], trend.get("keyword") or trend.get("category_name"),
            trend["start_rank"], trend["end_rank"], trend["direction"], trend["change_positions"],
            trend["valid_days"], trend["summary"],
        ])

    base_headers = [
        "商品", "ASIN", "标题", "Highlights", "About this item", "产品图数量", "产品图链接", "上架日期",
        "品牌", "当前价格", "企业价", "划线价", "Coupon", "Deal",
        "大类BSR", "小类BSR", "Rating数量", "星级", "各星级评价占比", "Amazon's Choice", "Best Seller",
        "高退货率标签", "商品状态", "库存", "Featured Offer卖家", "发货方", "跟卖数量", "变体数量", "采集时间",
    ]
    keyword_headers = []
    keyword_data: dict[str, dict[str, tuple[int | None, int | None]]] = {}
    for item in keywords:
        keyword_headers.extend([f"{item.keyword} 自然位", f"{item.keyword} 广告位"])
        rows = _latest_keyword_rows(db, project_id, item.keyword)
        for asin in target_asins:
            matching = [row for row in rows if row["asin"] == asin]
            organic = min((row["organic_rank"] for row in matching if row["organic_rank"] is not None), default=None)
            ad = min((row["ad_rank"] for row in matching if row["ad_rank"] is not None), default=None)
            keyword_data.setdefault(asin, {})[item.keyword] = (organic, ad)
    current_sheet.append(base_headers + keyword_headers)

    for asin in target_asins:
        row = products.get(asin)
        if not row:
            current_sheet.append([identities.get(asin, asin), asin, "尚无成功采集数据"])
            continue
        values = [
            identities.get(asin, asin), asin, row["title"], row.get("highlights_text"), _format_bullets(row.get("about_items_json")),
            row.get("product_image_count"), _format_images(row.get("product_images_json")), row.get("date_first_available"), row["brand"],
            row["current_price"], row.get("business_price_text"), row["list_price"], row["coupon_text"], row["deal_text"],
            _format_main_bsr(row["category_ranks_json"], row["main_bsr"]),
            _format_bsr(row["category_ranks_json"], row["main_bsr"]), row["rating_count"],
            row["rating_value"], format_rating_breakdown(row.get("rating_breakdown_json")), row["amazon_choice"], row["best_seller"],
            _format_high_return_rate(row.get("high_return_rate")), _format_listing_status(row.get("listing_status")), row["availability"],
            row["featured_seller"], row.get("ships_from"), row["offer_count"], row["variation_count"], row["collected_at"],
        ]
        for item in keywords:
            organic, ad = keyword_data.get(asin, {}).get(item.keyword, (None, None))
            values.extend([organic if organic is not None else f"未进入前{item.pages}页", ad if ad is not None else "未观察到"])
        current_sheet.append(values)

    history_sheet.append(base_headers)
    for row in db.fetchall(
        """SELECT * FROM product_snapshots WHERE project_id=?
           AND (success=1 OR listing_status IN ('dog','removed')) ORDER BY collected_at, asin""",
        (project_id,),
    ):
        if row["asin"] not in target_asins:
            continue
        history_sheet.append([
            identities.get(row["asin"], row["asin"]), row["asin"], row["title"], row.get("highlights_text"), _format_bullets(row.get("about_items_json")),
            row.get("product_image_count"), _format_images(row.get("product_images_json")), row.get("date_first_available"), row["brand"],
            row["current_price"], row.get("business_price_text"), row["list_price"], row["coupon_text"], row["deal_text"],
            _format_main_bsr(row["category_ranks_json"], row["main_bsr"]),
            _format_bsr(row["category_ranks_json"], row["main_bsr"]), row["rating_count"],
            row["rating_value"], format_rating_breakdown(row.get("rating_breakdown_json")), row["amazon_choice"], row["best_seller"],
            _format_high_return_rate(row.get("high_return_rate")), _format_listing_status(row.get("listing_status")), row["availability"],
            row["featured_seller"], row.get("ships_from"), row["offer_count"], row["variation_count"], row["collected_at"],
        ])

    keyword_history_sheet.append(["采集时间", "商品", "关键词", "ASIN", "自然位", "广告位"])
    for item in keywords:
        rows = db.fetchall("SELECT * FROM search_snapshots WHERE project_id=? AND keyword=? ORDER BY collected_at, asin", (project_id, item.keyword))
        grouped: dict[tuple[str, str], list[dict]] = {}
        for row in rows:
            if row["asin"] in target_asins:
                grouped.setdefault((row["run_id"], row["asin"]), []).append(row)
        for (_, asin), matches in grouped.items():
            collected_at = max(row["collected_at"] for row in matches)
            organic = min((row["organic_rank"] for row in matches if row["organic_rank"] is not None), default=None)
            ad = min((row["ad_rank"] for row in matches if row["ad_rank"] is not None), default=None)
            keyword_history_sheet.append([collected_at, identities.get(asin, asin), item.keyword, asin, organic, ad])

    bestseller_sheet.append(["类目", "排名", "商品", "ASIN", "标题", "品牌", "价格", "Rating数量", "星级", "采集时间"])
    latest_bsr_run = db.fetchall("SELECT run_id FROM bestseller_snapshots WHERE project_id=? ORDER BY collected_at DESC LIMIT 1", (project_id,))
    if latest_bsr_run:
        for row in db.fetchall("SELECT * FROM bestseller_snapshots WHERE project_id=? AND run_id=? ORDER BY category_name, rank", (project_id, latest_bsr_run[0]["run_id"])):
            bestseller_sheet.append([row["category_name"], row["rank"], identities.get(row["asin"], row["asin"]), row["asin"], row["title"], row["brand"], row["price"], row["rating_count"], row["rating_value"], row["collected_at"]])

    new_asin_sheet.append(["首次发现", "商品", "类目", "当前排名", "ASIN", "上架日期", "日期来源", "上架天数", "筛选状态", "判断来源", "判断依据", "商品链接", "已提醒时间"])
    status_labels = {"pending": "待确认", "same": "同类竞品", "not_same": "非同类产品", "too_old": "超过3个月"}
    source_labels = {"auto": "规则判断", "manual": "人工确认"}
    date_source_labels = {"auto": "自动采集", "manual": "人工填写"}
    for row in db.fetchall("SELECT * FROM bsr_new_candidates WHERE project_id=? ORDER BY first_seen_at DESC", (project_id,)):
        new_asin_sheet.append([
            row["first_seen_at"], identities.get(row["asin"], row["asin"]), row["category_name"], row["current_rank"],
            row["asin"], row["date_first_available"], date_source_labels.get(row.get("date_source"), ""),
            row["age_days"], status_labels.get(row["relevance_status"], row["relevance_status"]),
            source_labels.get(row["classification_source"], row["classification_source"]), row["relevance_reason"], row["detail_url"], row["alerted_at"],
        ])
        if row["detail_url"]:
            link_cell = new_asin_sheet.cell(row=new_asin_sheet.max_row, column=12)
            link_cell.hyperlink = row["detail_url"]
            link_cell.style = "Hyperlink"

    error_sheet.append(["任务", "目标", "错误类型", "错误信息", "重试次数", "截图", "HTML", "时间"])
    for row in db.fetchall("SELECT * FROM collection_errors WHERE project_id=? AND run_id=? ORDER BY created_at", (project_id, run_id)):
        error_sheet.append([row["task_type"], row["target"], row["error_type"], row["error_message"], row["retry_count"], row["screenshot_path"], row["html_path"], row["created_at"]])

    run_sheet.append(["批次ID", "任务", "开始时间", "结束时间", "状态", "任务数", "成功", "失败", "验证码"])
    for row in db.fetchall("SELECT * FROM collection_runs WHERE project_id=? ORDER BY started_at", (project_id,)):
        run_sheet.append([row["run_id"], row["task_type"], row["started_at"], row["finished_at"], row["status"], row["total_tasks"], row["success_count"], row["failed_count"], row["captcha_count"]])

    for sheet in workbook.worksheets:
        _style(sheet)
    for hidden_name in ("今日汇总", "今日关键词汇总", "竞品历史", "关键词排名历史", "BSR前100", "运行记录"):
        workbook[hidden_name].sheet_state = "hidden"
    project_dir = REPORT_DIR / project.project_name
    project_dir.mkdir(parents=True, exist_ok=True)
    main_report = project_dir / report_filename(project.project_name)
    change_report = project_dir / legacy_change_report_filename(project.project_name)
    workbook.save(main_report)
    # The former change-only workbook duplicated the first sheet of the main
    # report. Remove that rebuildable legacy artifact now that summaries have
    # dedicated daily/weekly downloads.
    change_report.unlink(missing_ok=True)
    return str(main_report)
