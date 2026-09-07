from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import load_projects
from app.notifications.serverchan import _daily_change_items
from app.presentation import (
    FIELD_LABELS, SOURCE_LABELS, build_product_identities, format_change_summary,
)
from app.storage.database import Database


HEADERS = [
    "日期", "项目", "商品", "变化项目", "变化摘要", "原值", "新值", "采集时间",
    "ASIN", "商品链接", "来源", "关键词", "类目",
]


def _style(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, column in enumerate(sheet.columns, 1):
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
        sheet.column_dimensions[get_column_letter(index)].width = max(width, 10)


def _add_link(cell, value: str | None) -> None:
    if value and str(value).startswith(("https://", "http://")):
        cell.hyperlink = str(value)
        cell.style = "Hyperlink"


def build_daily_excel_report(report_date: date, db: Database | None = None) -> BytesIO:
    database = db or Database()
    projects = {item.project_id: item.project_name for item in load_projects()}
    identities = {
        project_id: build_product_identities(database, project_id)
        for project_id in projects
    }
    changes = _daily_change_items(database, report_date)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "昨日日报"
    sheet.append(HEADERS)
    if not changes:
        sheet.append([report_date.isoformat(), "全部项目", "昨日未发现竞品变化"])
    for row in changes:
        project_id = str(row.get("project_id") or "")
        asin = str(row.get("asin") or "")
        product_url = f"https://www.amazon.com/dp/{asin}" if asin else ""
        old_value = row.get("old_value")
        new_value = row.get("new_value")
        sheet.append([
            report_date.isoformat(), projects.get(project_id, project_id),
            identities.get(project_id, {}).get(asin, asin or "—"),
            FIELD_LABELS.get(row.get("field_name"), row.get("field_name") or "变化"),
            format_change_summary(row), old_value, new_value, row.get("event_time"),
            asin, product_url, SOURCE_LABELS.get(row.get("source_type"), row.get("source_type")),
            row.get("keyword"), row.get("category_name"),
        ])
        _add_link(sheet.cell(sheet.max_row, 6), old_value)
        _add_link(sheet.cell(sheet.max_row, 7), new_value)
        _add_link(sheet.cell(sheet.max_row, 10), product_url)
    _style(sheet)

    payload = BytesIO()
    workbook.save(payload)
    payload.seek(0)
    return payload
