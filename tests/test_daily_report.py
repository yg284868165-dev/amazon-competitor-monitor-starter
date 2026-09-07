from datetime import date

from openpyxl import load_workbook

import app.reports.daily_report as daily_report
from app.config import Project
from app.reports.daily_report import build_daily_excel_report
from app.storage.database import Database


def test_daily_excel_contains_changes_and_clickable_asin_link(monkeypatch, tmp_path):
    db = Database(tmp_path / "test.db")
    db.insert("change_events", {
        "run_id": "run", "project_id": "project", "event_type": "field_changed",
        "severity": "medium", "source_type": "product", "asin": "B000000001",
        "field_name": "image_hash", "old_value": "https://example.com/old.jpg",
        "new_value": "https://example.com/new.jpg", "event_time": "2026-09-02T14:00:00",
        "message": "主图变化",
    })
    monkeypatch.setattr(
        daily_report, "load_projects",
        lambda: [Project("project", "测试项目", "90001")],
    )
    monkeypatch.setattr(
        daily_report, "build_product_identities",
        lambda *_: {"B000000001": "测试商品｜B000000001"},
    )

    workbook = load_workbook(build_daily_excel_report(date(2026, 9, 2), db))
    sheet = workbook["昨日日报"]
    headers = [cell.value for cell in sheet[1]]
    values = [cell.value for cell in sheet[2]]

    assert values[headers.index("项目")] == "测试项目"
    assert values[headers.index("变化摘要")] == "主图变化"
    assert values[headers.index("采集时间")] == "2026-09-02T14:00:00"
    assert "发生时间" not in headers
    assert sheet.cell(2, headers.index("ASIN") + 1).value == "B000000001"
    assert sheet.cell(2, headers.index("商品链接") + 1).hyperlink.target == "https://www.amazon.com/dp/B000000001"
    assert sheet.cell(2, headers.index("原值") + 1).hyperlink.target == "https://example.com/old.jpg"
    assert sheet.cell(2, headers.index("新值") + 1).hyperlink.target == "https://example.com/new.jpg"
