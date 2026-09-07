from datetime import date

from openpyxl import load_workbook

import app.reports.weekly_report as weekly_report
from app.config import Project
from app.reports.weekly_report import build_weekly_excel_report
from app.storage.database import Database


def test_weekly_excel_contains_rating_and_trend_rows_with_links(monkeypatch, tmp_path):
    db = Database(tmp_path / "test.db")
    for index, (old, new, timestamp) in enumerate((
        (100, 103, "2026-08-31T08:00:00"),
        (103, 108, "2026-09-06T20:00:00"),
    )):
        db.insert("change_events", {
            "run_id": f"run-{index}", "project_id": "project",
            "event_type": "field_changed", "severity": "low",
            "source_type": "product", "asin": "B000000001",
            "field_name": "rating_count", "old_value": str(old), "new_value": str(new),
            "event_time": timestamp, "message": "评价变化",
        })
    db.insert("change_events", {
        "run_id": "run-price", "project_id": "project",
        "event_type": "field_changed", "severity": "high",
        "source_type": "product", "asin": "B000000001",
        "field_name": "current_price", "old_value": "10", "new_value": "8",
        "event_time": "2026-09-02T14:00:00", "message": "价格变化",
    })
    monkeypatch.setattr(
        weekly_report, "load_projects",
        lambda: [Project("project", "测试项目", "90001")],
    )
    monkeypatch.setattr(
        weekly_report, "build_product_identities",
        lambda *_: {"B000000001": "测试商品｜B000000001"},
    )
    monkeypatch.setattr(weekly_report, "analyze_all_weekly_trends", lambda *_: [{
        "project_id": "project", "metric": "main_bsr", "asin": "B000000001",
        "product": "测试商品｜B000000001", "category_name": "Health & Household",
        "summaries": ["7日从20000名提升到16000名", "14日数据不足", "28日数据不足"],
        "presence_change": None,
    }])

    workbook = load_workbook(build_weekly_excel_report(date(2026, 9, 6), db))
    recap_sheet = workbook["本周重点复盘"]
    recap_headers = [cell.value for cell in recap_sheet[1]]
    trend_sheet = workbook["评价与排名趋势"]
    trend_headers = [cell.value for cell in trend_sheet[1]]

    assert recap_sheet.max_row == 2
    assert recap_sheet.cell(2, recap_headers.index("变化摘要") + 1).value == "当前价格：$10 → $8"
    assert recap_sheet.cell(2, recap_headers.index("截至周末状态") + 1).value == "截至周末仍保持变化"
    assert recap_sheet.cell(2, recap_headers.index("商品链接") + 1).hyperlink.target == "https://www.amazon.com/dp/B000000001"
    assert trend_sheet.max_row == 3
    assert trend_sheet.cell(2, trend_headers.index("近1周") + 1).value == "100 → 108（增加8条）"
    assert trend_sheet.cell(3, trend_headers.index("周报项目") + 1).value == "大类目BSR"
    assert trend_sheet.cell(3, trend_headers.index("近2周") + 1).value == "14日数据不足"


def test_weekly_excel_contains_compact_ad_insight_and_raw_daypart_evidence(monkeypatch, tmp_path):
    db = Database(tmp_path / "test.db")
    observation = {
        "run_id": "run-ad", "run_type": "all", "project_id": "project", "project_name": "测试项目",
        "asin": "B000000001", "product": "测试商品｜B000000001", "keyword": "cleaning stone",
        "status": "observed", "ad_rank": 4, "china_datetime": "2026-09-02T20:03+08:00",
        "pacific_datetime": "2026-09-02T05:03-07:00", "pacific_date": "2026-09-02",
        "slot": "05:00", "slot_order": 5,
    }
    monkeypatch.setattr(weekly_report, "load_projects", lambda: [Project("project", "测试项目", "90001")])
    monkeypatch.setattr(weekly_report, "build_product_identities", lambda *_: {"B000000001": "测试商品｜B000000001"})
    monkeypatch.setattr(weekly_report, "collect_ad_observations", lambda *_: [observation])
    monkeypatch.setattr(weekly_report, "summarize_ad_insights", lambda *_: [{
        "project_id": "project", "project_name": "测试项目", "product": "测试商品｜B000000001",
        "asin": "B000000001", "keyword": "cleaning stone", "observation": "重复分时规律",
        "inference": "疑似分时竞价", "confidence": "中",
        "slot_statistics": {"05:00": {"observed": 1, "successful": 1, "median_rank": 4}},
    }])

    workbook = load_workbook(build_weekly_excel_report(date(2026, 9, 6), db))

    assert workbook["广告投放观察"]["F2"].value == "重复分时规律"
    assert "05:03 第4位" in workbook["广告分时轨迹"]["G2"].value
    assert workbook["广告采集明细"]["H2"].value == 4
    assert workbook["广告采集明细"]["L2"].hyperlink.target == "https://www.amazon.com/dp/B000000001"
