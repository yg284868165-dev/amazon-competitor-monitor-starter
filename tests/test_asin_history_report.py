from io import BytesIO

from openpyxl import load_workbook

import app.reports.asin_history_report as report
from app.config import Project
from app.storage.database import Database


def test_action_history_excel_contains_evidence_sheets(monkeypatch, tmp_path):
    data = {
        "amazon_url": "https://www.amazon.com/dp/B000000001",
        "identity": "内部名称｜Brand｜B000000001",
        "asin": "B000000001",
        "range": {
            "from": "2026-08-17T12:00:00", "to": "2026-09-16T12:00:00",
            "sample_count": 2,
        },
        "summary": {
            "important_nodes": 1, "price_promo": 1, "listing": 1,
            "operations": 0, "platform": 0,
        },
        "actions": [{
            "observed_at": "2026-09-10T08:01:00", "title": "价格与促销、Listing内容组合变化",
            "category_labels": ["价格与促销", "Listing内容"], "run_id": "action",
            "items": [{
                "label": "主图", "summary": "主图变化", "old": ["https://example.com/old.jpg"],
                "new": ["https://example.com/new.jpg"], "display_type": "images",
            }],
            "impacts": [{
                "metric_label": "小类目BSR", "context": "Cleaning Tools", "run_id": "action",
                "points": {
                    "before": {"status": "ranked", "value": 42, "observed_at": "2026-09-09T08:00:00"},
                    "day_1": {"status": "ranked", "value": 31, "observed_at": "2026-09-11T08:00:00"},
                    "day_3": {"status": "ranked", "value": 20, "observed_at": "2026-09-13T08:00:00"},
                    "day_7": {"status": "pending", "value": None, "observed_at": None},
                },
            }],
        }],
        "listing_changes": [{
            "observed_at": "2026-09-10T08:01:00", "run_id": "action",
            "item": {
                "label": "主图", "summary": "主图变化", "old": ["https://example.com/old.jpg"],
                "new": ["https://example.com/new.jpg"], "display_type": "images",
            },
        }],
        "impacts": [],
        "intraday": {
            "product": [{
                "collected_at": "2026-09-10T08:00:00", "price": 10.99, "coupon": None,
                "deal": None, "business_price": None, "availability": "In Stock", "offer_count": 1,
                "featured_seller": "Brand", "ships_from": "Amazon", "listing_status": "active",
            }],
            "ads": [{
                "collected_at": "2026-09-10T08:05:00", "keyword": "cleaning stone",
                "status": "not_observed", "ad_rank": None,
            }],
        },
    }
    monkeypatch.setattr(report, "build_asin_history", lambda db, project_id, asin, days: data)
    monkeypatch.setattr(report, "load_projects", lambda: [Project("project", "测试项目", "90001")])

    payload = report.build_asin_history_excel(
        "project", "B000000001", 30, Database(tmp_path / "test.db"),
    )
    assert isinstance(payload, BytesIO)
    workbook = load_workbook(payload, data_only=True)
    assert workbook.sheetnames == [
        "动作摘要", "动作时间线", "Listing版本对比",
        "日内商品状态", "广告分时证据",
    ]
    assert workbook["动作时间线"]["D3"].value == "后续排名对照"
    assert workbook["动作时间线"]["G3"].value.startswith("第42名")
    assert workbook["Listing版本对比"]["C2"].hyperlink.target == "https://example.com/old.jpg"
    assert workbook["广告分时证据"]["C2"].value == "未观察到广告"
