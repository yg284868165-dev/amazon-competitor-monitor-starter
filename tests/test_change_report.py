from openpyxl import Workbook

from app.reports.excel_report import _append_changes


def test_change_rows_are_chinese_and_include_clickable_product_link():
    sheet = Workbook().active
    _append_changes(sheet, [{
        "severity": "high", "source_type": "product", "asin": "B000000010",
        "keyword": None, "category_name": None, "field_name": "current_price",
        "old_value": "10", "new_value": "9", "change_value": -1,
        "message": "价格变化", "event_time": "2026-08-31T10:00:00",
    }], {"B000000010": "主竞品｜ExampleCo｜B000000010"})
    assert sheet.cell(1, 1).value == "当前价格"
    assert sheet.cell(1, 2).value == "主竞品｜ExampleCo｜B000000010"
    assert sheet.cell(1, 3).value == "当前价格：$10 → $9"
    assert sheet.cell(1, 8).value == "商品详情"
    assert sheet.cell(1, 10).value == "https://www.amazon.com/dp/B000000010"
    assert sheet.cell(1, 10).hyperlink.target == "https://www.amazon.com/dp/B000000010"
